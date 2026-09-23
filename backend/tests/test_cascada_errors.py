"""CASCADA error data: the same call as rebuffering with one metric changed, and a count.

`tests/fixtures/cascada_errors_response.json` is a real `target_metrics[]=error_count` response
with its two series trimmed to six hours each and its signed download link replaced. Its
`reference_time` block, its `unit` map and its row shape are untouched — including the one
thing that differs from what was asked for: every row carries its value under `errors`, not
`error_count`.

The API tests point `app.cascada.client` at the fixture server, which answers the rebuffering
call and the error call from their own saved responses, so each test proves which call went
out as well as what came back.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
from collections.abc import Iterator
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlsplit

import pytest
from fastapi.testclient import TestClient

from app.cascada import client, errors, exports, series
from app.config import Thresholds, get_settings, get_thresholds
from app.main import create_app
from tests.fixtures.server import FixtureServer, Route

FIXTURES = Path(__file__).parent / "fixtures"
ERRORS_FIXTURE = FIXTURES / "cascada_errors_response.json"
REBUFFERING_FIXTURE = FIXTURES / "cascada_response.json"

REALTIME_PATH = "/api/data/v1/realtime"
SESSION = "Cookie: sessionid=test-session-value; csrftoken=test-csrf"

T = Thresholds()

ORIGIN_START = dt.datetime(2026, 9, 21, 0, 0, tzinfo=dt.UTC)
ORIGIN_END = dt.datetime(2026, 9, 21, 5, 59, tzinfo=dt.UTC)

# What the six saved origin hours add up to, computed once from the fixture by hand.
FIXTURE_MINUTES = 360
FIXTURE_AVERAGE = 643.0
FIXTURE_MAX = 961.0
FIXTURE_TOTAL = 231480.0


def payload() -> dict[str, Any]:
    return dict(json.loads(ERRORS_FIXTURE.read_text(encoding="utf-8")))


def fixture_window() -> series.Window:
    return series.Window(start=ORIGIN_START, end=ORIGIN_END)


def parsed() -> series.ChannelSeries:
    return series.parse(payload(), requested=fixture_window(), metric=series.ERRORS)


def with_threshold(value: float) -> Thresholds:
    return Thresholds(cascada_error_threshold_per_min=value)


# -- the call ----------------------------------------------------------------


def test_the_error_call_differs_from_the_rebuffering_call_only_in_its_metric() -> None:
    span = fixture_window()
    common = {"channel_name": "Samsung Television Network", "channel_id": "US300068X7"}
    rebuffering = parse_qs(urlsplit(client.build_url(window=span, limit=10, **common)).query)
    errors_query = parse_qs(
        urlsplit(client.build_url(window=span, limit=10, metric=series.ERRORS, **common)).query
    )

    assert rebuffering.pop("target_metrics[]") == ["rebuffering_ratio"]
    assert errors_query.pop("target_metrics[]") == ["error_count"]
    assert rebuffering == errors_query


def test_the_error_url_matches_the_one_the_dashboard_sends() -> None:
    """The query string, parameter for parameter and in order, as CASCADA was shown it."""
    span = series.Window(
        start=dt.datetime.fromtimestamp(1789948800, dt.UTC),
        end=dt.datetime.fromtimestamp(1790139840, dt.UTC),
    )
    url = client.build_url(
        channel_name="Samsung Television Network",
        channel_id="US300068X7",
        window=span,
        limit=13264,
        base_url="https://cascada.samsungcloud.tv",
        metric=series.ERRORS,
    )
    assert url == (
        "https://cascada.samsungcloud.tv/api/data/v1/realtime?channel_group_name=Master+Group+-"
        "+Local+Channel+Included&channel_name=Samsung+Television+Network&channel_id=US300068X7"
        "&channel_country=ALL&platform=Samsung+TV&model=ALL&from=1789948800&to=1790139840"
        "&limit=13264&compare_with=1WEEK&target_metrics%5B%5D=error_count&export_to_gcs=true"
    )


# -- the response ------------------------------------------------------------


def test_the_value_is_read_from_errors_although_error_count_was_asked_for() -> None:
    """Reading the key that was asked for would draw a week of gaps on a measured channel."""
    result = parsed()

    assert len(result.origin) == FIXTURE_MINUTES
    assert len(result.comparison) == FIXTURE_MINUTES
    assert all(point.value is not None for point in result.origin)
    assert result.origin[0].value == 881.0
    assert result.unit == "times"


def test_a_row_keyed_error_count_is_read_as_well() -> None:
    rows = [
        {"target_time": "2026-09-21T00:00:00+00:00", "date_category": "origin", "error_count": 7}
    ]
    result = series.parse({"data": rows}, requested=fixture_window(), metric=series.ERRORS)
    assert result.origin[0].value == 7.0


def test_a_response_carrying_another_metric_is_refused_naming_what_it_carried() -> None:
    """A rebuffering answer to an error call is an error, not a window with nothing measured."""
    rebuffering = json.loads(REBUFFERING_FIXTURE.read_text(encoding="utf-8"))

    with pytest.raises(series.CascadaParseError) as caught:
        series.parse(rebuffering, requested=fixture_window(), metric=series.ERRORS)

    message = str(caught.value)
    assert "`errors`" in message and "`error_count`" in message
    assert "rebuffering_ratio" in message


def test_an_errors_response_does_not_satisfy_the_rebuffering_parser_either() -> None:
    with pytest.raises(series.CascadaParseError):
        series.parse(payload(), requested=fixture_window())


def test_a_null_error_minute_is_a_gap_and_not_a_zero() -> None:
    rows = [
        {"target_time": "2026-09-21T00:00:00+00:00", "date_category": "origin", "errors": 10.0},
        {"target_time": "2026-09-21T00:01:00+00:00", "date_category": "origin", "errors": None},
    ]
    result = series.parse({"data": rows}, requested=fixture_window(), metric=series.ERRORS)
    stats = errors.summarise(result.origin, result.comparison, T)

    assert stats.average_per_min == pytest.approx(10.0)
    assert stats.total == pytest.approx(10.0)
    assert stats.minutes_counted == 1
    assert stats.minutes_missing == 1


# -- the figures -------------------------------------------------------------


def test_every_figure_comes_from_the_current_week_alone() -> None:
    result = parsed()
    stats = errors.summarise(result.origin, result.comparison, T)

    assert stats.average_per_min == pytest.approx(FIXTURE_AVERAGE)
    assert stats.max_per_min == pytest.approx(FIXTURE_MAX)
    assert stats.total == pytest.approx(FIXTURE_TOTAL)
    assert stats.minutes_counted == FIXTURE_MINUTES
    # Last week is reported, and only as the comparison.
    comparison = [p.value for p in result.comparison if p.value is not None]
    assert stats.previous_week_average_per_min == pytest.approx(sum(comparison) / len(comparison))
    assert stats.week_over_week_delta_per_min == pytest.approx(
        FIXTURE_AVERAGE - sum(comparison) / len(comparison)
    )
    assert stats.week_over_week_change_pct is not None and stats.week_over_week_change_pct < 0


def test_with_no_threshold_set_no_minute_is_marked_and_no_verdict_is_stated() -> None:
    """A count scales with the audience, so the default states nothing it cannot back."""
    result = parsed()
    stats = errors.summarise(result.origin, result.comparison, T)

    assert T.cascada_error_threshold_per_min == 0.0
    assert stats.threshold_per_min is None
    assert stats.minutes_above == 0
    assert stats.percent_time_above is None
    assert stats.above_threshold is False


def test_a_threshold_marks_the_minutes_above_it_and_judges_the_channel_on_its_average() -> None:
    result = parsed()

    below = errors.summarise(result.origin, result.comparison, with_threshold(900))
    assert below.threshold_per_min == 900
    assert below.minutes_above > 0
    # Some minutes breach 900 but the average of 643 does not.
    assert below.above_threshold is False

    above = errors.summarise(result.origin, result.comparison, with_threshold(500))
    assert above.above_threshold is True


# -- the report --------------------------------------------------------------


def window_entry(thresholds: Thresholds = T) -> errors.ErrorWindow:
    result = parsed()
    return errors.to_window(
        service_id="US300068X7",
        channel_name="Samsung Television Network",
        country="US",
        series=result,
        window=fixture_window(),
        thresholds=thresholds,
    )


def test_the_error_csv_is_a_table_of_counts_with_its_context_after_it() -> None:
    text = exports.errors_csv(window_entry())
    rows = list(csv.reader(io.StringIO(text)))

    assert rows[0] == list(exports.ERROR_COLUMNS)
    assert rows[1] == ["2026-09-21T00:00:00+00:00", "881", ""]
    about = {row[0]: row[1] for row in rows[FIXTURE_MINUTES + 2 :] if len(row) == 2}
    assert about["report"] == "Samsung TV Plus — CASCADA error report"
    assert about["total_errors"] == "231480"
    assert about["threshold_per_min"].startswith("not set")


def test_the_error_csv_marks_minutes_once_a_threshold_is_set() -> None:
    rows = list(csv.reader(io.StringIO(exports.errors_csv(window_entry(with_threshold(900))))))
    marks = {row[2] for row in rows[1 : FIXTURE_MINUTES + 1]}
    assert marks == {"yes", "no"}


def test_the_error_workbook_has_the_minutes_and_the_context() -> None:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(exports.errors_xlsx(window_entry())))
    assert book.sheetnames == ["minutes", exports.ABOUT_SHEET]
    assert book["minutes"].max_row == FIXTURE_MINUTES + 1


def test_the_error_report_filename_says_errors() -> None:
    now = dt.datetime(2026, 9, 23, tzinfo=dt.UTC)
    assert exports.channel_filename("US300068X7", "csv", now, prefix="errors") == (
        "errors_US300068X7_20260923.csv"
    )
    # The rebuffering name is unchanged.
    assert (
        exports.channel_filename("US300068X7", "csv", now) == "rebuffering_US300068X7_20260923.csv"
    )


# -- the endpoints -----------------------------------------------------------


def _answer(query: dict[str, list[str]]) -> bytes:
    """Each metric's saved response, chosen by the metric the call asked for."""
    metric = query.get("target_metrics[]", [""])[0]
    source = ERRORS_FIXTURE if metric == "error_count" else REBUFFERING_FIXTURE
    return source.read_bytes()


@pytest.fixture()
def cascada(origin: FixtureServer) -> Iterator[FixtureServer]:
    settings = get_settings()
    previous = settings.cascada_base_url
    settings.cascada_base_url = origin.base_url
    origin.add(REALTIME_PATH, Route(body=_answer, content_type="application/json"))
    try:
        yield origin
    finally:
        settings.cascada_base_url = previous


@pytest.fixture()
def api(cascada: FixtureServer) -> Iterator[TestClient]:
    with TestClient(create_app()) as test_client:
        assert test_client.delete("/api/cascada/session").status_code == 200
        response = test_client.put("/api/cascada/session", json={"cookie_header": SESSION})
        assert response.status_code == 200, response.text
        try:
            yield test_client
        finally:
            test_client.delete("/api/cascada/session")


def channel(service_id: str) -> dict[str, str]:
    return {"service_id": service_id, "channel_name": "Samsung Television Network", "country": "US"}


def sent_metrics(server: FixtureServer) -> list[str]:
    return [
        parse_qs(entry["query"]).get("target_metrics[]", [""])[0]
        for entry in server.request_log
        if entry["path"] == REALTIME_PATH
    ]


def test_the_error_endpoint_asks_for_error_count_and_returns_the_counts(
    api: TestClient, cascada: FixtureServer
) -> None:
    cascada.clear_log()
    response = api.get(
        "/api/cascada/channel/errors", params={**channel("US-ERR-1"), "refresh": "true"}
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert sent_metrics(cascada) == ["error_count"]
    assert body["metric"] == "errors"
    assert body["unit"] == "times"
    assert body["total"] == pytest.approx(FIXTURE_TOTAL)
    assert body["average_per_min"] == pytest.approx(FIXTURE_AVERAGE)
    assert body["threshold_per_min"] is None
    assert len(body["origin"]) == FIXTURE_MINUTES
    assert len(body["comparison"]) == FIXTURE_MINUTES
    # The session never comes back to a page.
    assert "test-session-value" not in response.text


def test_a_reopened_error_window_is_served_from_the_store(
    api: TestClient, cascada: FixtureServer
) -> None:
    first = api.get(
        "/api/cascada/channel/errors", params={**channel("US-ERR-2"), "refresh": "true"}
    )
    assert first.status_code == 200
    cascada.clear_log()

    second = api.get("/api/cascada/channel/errors", params=channel("US-ERR-2"))

    assert second.status_code == 200
    assert second.json()["cached"] is True
    assert sent_metrics(cascada) == []


def test_the_error_and_rebuffering_windows_of_one_channel_are_stored_apart(
    api: TestClient, cascada: FixtureServer
) -> None:
    """One channel, one window, two metrics: neither cache entry may answer for the other."""
    ident = channel("US-ERR-3")
    assert (
        api.get("/api/cascada/channel/errors", params={**ident, "refresh": "true"}).status_code
        == 200
    )
    cascada.clear_log()

    rebuffering = api.get("/api/cascada/channel", params=ident)

    assert rebuffering.status_code == 200, rebuffering.text
    assert "average_pct" in rebuffering.json()
    assert sent_metrics(cascada) == ["rebuffering_ratio"]


def test_a_threshold_changed_in_settings_applies_to_a_stored_error_window(
    api: TestClient,
) -> None:
    ident = channel("US-ERR-4")
    assert (
        api.get("/api/cascada/channel/errors", params={**ident, "refresh": "true"}).status_code
        == 200
    )
    try:
        get_thresholds().cascada_error_threshold_per_min = 500
        body = api.get("/api/cascada/channel/errors", params=ident).json()
        assert body["cached"] is True
        assert body["threshold_per_min"] == 500
        assert body["above_threshold"] is True
    finally:
        get_thresholds().cascada_error_threshold_per_min = 0.0


def test_the_error_report_downloads_as_csv_and_xlsx(api: TestClient) -> None:
    ident = channel("US-ERR-5")
    csv_response = api.get("/api/cascada/channel/errors/report.csv", params=ident)
    assert csv_response.status_code == 200, csv_response.text
    assert csv_response.headers["content-type"].startswith("text/csv")
    assert 'filename="errors_US-ERR-5_' in csv_response.headers["content-disposition"]
    assert csv_response.text.startswith(",".join(exports.ERROR_COLUMNS))

    xlsx_response = api.get("/api/cascada/channel/errors/report.xlsx", params=ident)
    assert xlsx_response.status_code == 200
    assert xlsx_response.content[:2] == b"PK"

    assert api.get("/api/cascada/channel/errors/report.pdf", params=ident).status_code == 400


def test_an_error_call_with_no_session_is_the_same_401_as_rebuffering(api: TestClient) -> None:
    assert api.delete("/api/cascada/session").status_code == 200
    response = api.get(
        "/api/cascada/channel/errors",
        params={"service_id": "US-ERR-NEVER", "channel_name": "Unknown", "country": "US"},
    )
    assert response.status_code == 401
    assert "Settings tab" in response.json()["detail"]
