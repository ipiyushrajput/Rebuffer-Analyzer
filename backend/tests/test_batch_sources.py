"""Automated batches with a data source: what is recorded, what is defaulted, what is stated.

A batch reads rebuffering from realtime or historical CASCADA data. The source is chosen by the
button, recorded with the batch, defaulted to historical for a scheduled run, kept by a re-run,
and stated beside every average the report prints. The correlation of a batch's spike windows
with aging reads the aging runs the *previous* batch started, because those are the runs that
watched the week this batch measured.
"""

from __future__ import annotations

import asyncio
import datetime as dt
import io
from typing import Any

import pytest
from fastapi.testclient import TestClient

from app.batch import exports, reporting, runner, store
from app.batch.settings import DEFAULTS, BatchSettings
from app.cascada import service as cascada
from app.cascada.series import Point, Stats, Window
from app.cascada.store import ChannelWindow
from app.db import session as db_session
from app.db.models import Batch, BatchItem, BatchLog
from app.main import create_app

WINDOW = Window(
    start=dt.datetime(2026, 9, 16, tzinfo=dt.UTC),
    end=dt.datetime(2026, 9, 22, tzinfo=dt.UTC),
)


async def _clear() -> None:
    from sqlalchemy import delete

    async with db_session.session_scope() as session:
        await session.execute(delete(BatchLog))
        await session.execute(delete(BatchItem))
        await session.execute(delete(Batch))


@pytest.fixture(autouse=True)
def api() -> Any:
    with TestClient(create_app()) as client:
        for task in list(runner._tasks.values()):
            task.cancel()
        runner._tasks.clear()
        runner._cancelled.clear()
        client.portal.call(_clear)  # type: ignore[attr-defined]
        yield client


def _daily(service_id: str, name: str, average: float, *, above: bool) -> ChannelWindow:
    days = [WINDOW.start + dt.timedelta(days=n) for n in range(7)]
    return ChannelWindow(
        service_id=service_id,
        channel_name=name,
        country="GB",
        window=WINDOW,
        stats=Stats(
            average_pct=average, max_pct=average, max_at=days[0], minutes_above=7 if above else 0,
            minutes_counted=7, minutes_missing=0, percent_time_above=100.0 if above else 0.0,
            previous_week_average_pct=None, threshold_pct=0.25, above_threshold=above,
        ),
        origin=[Point(at=day, value=average) for day in days],
        comparison=[],
        fetched_at=dt.datetime.now(dt.UTC),
        truncated=False,
        source="historical",
        granularity="day",
        details={
            "returned_dates": [d.date().isoformat() for d in days],
            "returned_days": {"first_day": "2026-09-16", "last_day": "2026-09-22",
                              "days": 7, "label": "2026-09-16 → 2026-09-22"},
            "days_with_data": 7, "days_expected": 7,
        },
    )  # fmt: skip


def _historical_scan() -> cascada.Scan:
    scan = cascada.Scan(
        id="hist",
        country="GB",
        window=WINDOW,
        total=4,
        threshold_pct=0.25,
        source="historical",
        urls={"GB1": "https://cdn.example/1.m3u8", "GB2": "https://cdn.example/2.m3u8",
              "GB3": "https://cdn.example/3.m3u8"},
    )  # fmt: skip
    scan.results = {
        "GB1": _daily("GB1", "Worst", 0.9, above=True),
        "GB2": _daily("GB2", "Clean", 0.1, above=False),
    }
    scan.insufficient = [
        {"service_id": "GB3", "channel_name": "Thin", "days_with_data": 2, "days_expected": 7,
         "average_pct": 0.8}
    ]  # fmt: skip
    scan.no_provider = [
        {"service_id": "GB4", "channel_name": "Orphan", "reason": "Provider not found: none."}
    ]
    scan.done = 4
    scan.status = "COMPLETED"
    scan.finished_at = dt.datetime.now(dt.UTC)
    return scan


# -- recording the source ----------------------------------------------------


def test_the_scheduled_default_is_historical() -> None:
    assert BatchSettings().scheduled_data_source == "historical"
    assert DEFAULTS.snapshot()["scheduled_data_source"] == "historical"


@pytest.mark.parametrize(
    ("kind", "source", "expected"),
    [
        (store.MANUAL, "historical", "historical"),
        (store.MANUAL, "realtime", "realtime"),
        # A caller that predates the choice starts what it always started.
        (store.MANUAL, None, "realtime"),
        # A scheduled run reads the setting, historical by default.
        (store.SCHEDULED, None, "historical"),
    ],
)
def test_the_source_is_chosen_and_recorded_with_the_batch(
    api: TestClient,
    monkeypatch: pytest.MonkeyPatch,
    kind: str,
    source: str | None,
    expected: str,
) -> None:
    async def idle(batch_id: str, resume: bool = False) -> None:
        return None

    monkeypatch.setattr(runner, "_run", idle)
    started = api.portal.call(lambda: runner.start("GB", kind=kind, source=source))  # type: ignore[attr-defined]

    assert started["data_source"] == expected
    stored = api.portal.call(store.read, started["id"])  # type: ignore[attr-defined]
    assert stored is not None
    assert stored["data_source"] == expected
    assert stored["settings"]["data_source"] == expected
    # Left queued, the next app start would resume it for real.
    api.portal.call(lambda: store.update(started["id"], status=store.COMPLETED))  # type: ignore[attr-defined]


def test_the_start_endpoint_takes_the_source_of_the_button(
    api: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    async def idle(batch_id: str, resume: bool = False) -> None:
        return None

    monkeypatch.setattr(runner, "_run", idle)
    body = api.post("/api/batch/batches", json={"country": "GB", "source": "historical"}).json()
    assert body["data_source"] == "historical"

    listing = api.get("/api/batch/batches").json()["batches"]
    assert listing[0]["data_source"] == "historical"

    # A re-run reads the same source again.
    api.portal.call(lambda: store.update(body["id"], status=store.COMPLETED))  # type: ignore[attr-defined]
    again = api.post(f"/api/batch/batches/{body['id']}/rerun").json()
    assert again["data_source"] == "historical"
    api.portal.call(lambda: store.update(again["id"], status=store.COMPLETED))  # type: ignore[attr-defined]

    assert (
        api.post("/api/batch/batches", json={"country": "GB", "source": "weekly"}).status_code
        == 422
    )


def test_a_batch_from_before_the_choice_reads_as_realtime(api: TestClient) -> None:
    api.portal.call(  # type: ignore[attr-defined]
        lambda: store.create(batch_id="old", country="GB", kind=store.MANUAL, snapshot={})
    )
    assert api.get("/api/batch/batches/old").json()["data_source"] == "realtime"
    api.portal.call(lambda: store.update("old", status=store.COMPLETED))  # type: ignore[attr-defined]


# -- the historical scan inside a batch --------------------------------------


@pytest.mark.asyncio()
async def test_a_historical_batch_selects_on_the_historical_average_and_names_the_rest(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    scan = _historical_scan()
    asked: list[str] = []

    async def fake_start(
        country: str, concurrency: int | None = None, source: str = "realtime"
    ) -> cascada.Scan:
        asked.append(source)
        return scan

    monkeypatch.setattr(cascada, "start_scan", fake_start)
    batch_id = "hist-batch"
    await store.create(
        batch_id=batch_id,
        country="GB",
        kind=store.MANUAL,
        snapshot={**DEFAULTS.snapshot(), "data_source": "historical"},
    )
    await runner._scan(batch_id, "GB", DEFAULTS, "historical")

    assert asked == ["historical"]
    items = {item["service_id"]: item for item in await store.items_for(batch_id)}
    assert items["GB1"]["status"] == "PENDING"
    assert "GB2" not in items, "a channel below the threshold is not selected"
    assert items["GB3"]["status"] == runner.INSUFFICIENT_DATA
    assert "2 of 7" in (items["GB3"]["error"] or "")
    assert items["GB4"]["status"] == runner.NO_PROVIDER

    row = await store.read(batch_id) or {}
    await store.update(batch_id, status=store.COMPLETED)
    assert row["channels_above"] == 1
    # The stated window runs to the end of the last returned day.
    assert row["window_to"].startswith("2026-09-22T23:59:59")


@pytest.mark.asyncio()
async def test_a_realtime_batch_calls_the_scan_exactly_as_before(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls: list[tuple[Any, ...]] = []
    scan = _historical_scan()
    scan.source = "realtime"
    scan.insufficient, scan.no_provider = [], []

    async def fake_start(*args: Any, **kwargs: Any) -> cascada.Scan:
        calls.append((args, kwargs))
        return scan

    monkeypatch.setattr(cascada, "start_scan", fake_start)
    await store.create(batch_id="rt", country="GB", kind=store.MANUAL, snapshot=DEFAULTS.snapshot())
    await runner._scan("rt", "GB", DEFAULTS)
    await store.update("rt", status=store.COMPLETED)
    assert calls == [(("GB",), {})]


# -- the report --------------------------------------------------------------


def _batch(source: str) -> dict[str, Any]:
    return {
        "id": "b1",
        "country": "GB",
        "kind": "scheduled",
        "status": "COMPLETED",
        "data_source": source,
        "window_from": "2026-09-16T00:00:00+00:00",
        "window_to": "2026-09-22T23:59:59+00:00",
        "settings": {"threshold_pct": 0.25},
        "items": [
            {"service_id": "GB1", "channel_name": "Worst", "country": "GB", "status": "ANALYSED",
             "average_pct": 0.9, "summary": "done", "correlation": {}},
            {"service_id": "GB3", "channel_name": "Thin", "country": "GB",
             "status": "INSUFFICIENT_DATA", "average_pct": 0.8, "error": "Insufficient data: 2 of 7."},
            {"service_id": "GB4", "channel_name": "Orphan", "country": "GB",
             "status": "NO_PROVIDER", "error": "Provider not found."},
        ],
    }  # fmt: skip


@pytest.mark.parametrize("source", ["historical", "realtime"])
def test_the_average_column_names_its_source_and_dates(source: str) -> None:
    text = exports.csv_bytes(_batch(source)).decode()
    heading = text.splitlines()[0]
    assert f"Avg Rebuffering Ratio ({source}, 2026-09-16 → 2026-09-22)" in heading


def test_the_header_block_states_how_channels_were_selected_and_spikes_correlated() -> None:
    about = dict(row[:2] for row in exports._about(_batch("historical")))
    assert about["data_source"] == "historical"
    assert about["selected_by"] == "historical daily average (2026-09-16 → 2026-09-22 UTC)"
    assert about["spike_correlation"] == "no per-minute window was read"


def test_channels_not_judged_have_their_own_section_and_are_not_called_failures() -> None:
    text = exports.csv_bytes(_batch("historical")).decode()
    assert "Channels not judged (insufficient data or no provider)" in text
    selected_section = text.split("Channels not judged")[0]
    assert "INSUFFICIENT_DATA" not in selected_section

    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(exports.xlsx_bytes(_batch("historical"))))
    assert exports.NOT_JUDGED_SHEET in book.sheetnames
    statuses = [row[3] for row in book[exports.NOT_JUDGED_SHEET].iter_rows(values_only=True)]
    assert statuses[1:] == ["INSUFFICIENT_DATA", "NO_PROVIDER"]


# -- the correlation reads the aging that watched the measured week ----------


def _minutes(values: list[float | None], start: dt.datetime) -> ChannelWindow:
    points = [Point(at=start + dt.timedelta(minutes=n), value=v) for n, v in enumerate(values)]
    return ChannelWindow(
        service_id="GB1", channel_name="Worst", country="GB",
        window=Window(start=start, end=points[-1].at),
        stats=Stats(average_pct=0.5, max_pct=0.9, max_at=start, minutes_above=2,
                    minutes_counted=len(values), minutes_missing=0, percent_time_above=1.0,
                    previous_week_average_pct=None, threshold_pct=0.25, above_threshold=True),
        origin=points, comparison=[], fetched_at=dt.datetime.now(dt.UTC), truncated=False,
    )  # fmt: skip


@pytest.mark.asyncio()
async def test_the_previous_batchs_aging_run_is_correlated_with_this_weeks_spikes(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    start = dt.datetime(2026, 9, 16, tzinfo=dt.UTC)
    # Two spike minutes at 10:00 and 10:01 on the 18th.
    values: list[float | None] = [0.1] * (3 * 1440)
    values[2 * 1440 + 600] = values[2 * 1440 + 601] = 0.9
    minutes = _minutes(values, start)

    async def fake_window(**kwargs: Any) -> ChannelWindow:
        return minutes

    async def fake_period(job_id: str) -> tuple[dt.datetime, dt.datetime]:
        return start + dt.timedelta(hours=4), start + dt.timedelta(days=2, hours=23)

    async def fake_evidence(
        job_id: str | None,
    ) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
        at = (start + dt.timedelta(days=2, hours=10)).isoformat()
        return [{"first_seen": at, "rule_id": "SEG-001", "title": "Segment 404"}], []

    monkeypatch.setattr(reporting.cascada_service, "channel_window", fake_window)
    monkeypatch.setattr(reporting, "_aging_period", fake_period)
    monkeypatch.setattr(reporting, "_job_evidence", fake_evidence)

    item = {"service_id": "GB1", "channel_name": "Worst", "country": "GB"}
    result = await reporting._correlation(
        item, WINDOW, DEFAULTS, 0.25, source="historical", aging_job_id="aging-1",
        aging_batch_id="previous",
    )  # fmt: skip

    assert result["spike_source"] == "realtime per-minute"
    assert result["aging_job_id"] == "aging-1"
    assert result["aging_batch_id"] == "previous"
    assert result["coverage"].startswith("The realtime per-minute window covers")
    assert result["spike_count"] == 1
    assert result["matched_count"] == 1


@pytest.mark.asyncio()
async def test_an_aging_run_older_than_the_per_minute_window_is_stated_not_correlated(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    start = dt.datetime(2026, 9, 16, tzinfo=dt.UTC)
    minutes = _minutes([0.9, 0.9, 0.1], start)

    async def fake_window(**kwargs: Any) -> ChannelWindow:
        return minutes

    async def fake_period(job_id: str) -> tuple[dt.datetime, dt.datetime]:
        return start - dt.timedelta(days=2), start + dt.timedelta(days=1)

    monkeypatch.setattr(reporting.cascada_service, "channel_window", fake_window)
    monkeypatch.setattr(reporting, "_aging_period", fake_period)

    result = await reporting._correlation(
        {"service_id": "GB1", "channel_name": "Worst", "country": "GB"}, WINDOW, DEFAULTS, 0.25,
        source="realtime", aging_job_id="aging-old",
    )  # fmt: skip
    assert result["coverage"].startswith("Not correlated")
    assert result["windows"] == []


@pytest.mark.asyncio()
async def test_the_report_reads_the_previous_batchs_aging_run_for_each_channel() -> None:
    await store.create(batch_id="week1", country="GB", kind=store.SCHEDULED, snapshot={})
    await store.add_items(
        "week1", [{"service_id": "GB1", "channel_name": "Worst", "country": "GB"}]
    )
    await store.update_item("week1", "GB1", aging_job_id="aging-from-week1")
    await asyncio.sleep(0.01)
    await store.create(batch_id="week2", country="GB", kind=store.SCHEDULED, snapshot={})

    for finished in ("week1", "week2"):
        await store.update(finished, status=store.COMPLETED)
    current = await store.read("week2") or {}
    previous_id, aging = await reporting._previous_aging(current)
    assert previous_id == "week1"
    assert aging == {"GB1": "aging-from-week1"}


def test_a_report_that_never_built_is_built_when_it_is_downloaded(api: TestClient) -> None:
    """A build that failed part way leaves analysed channels with no summary; the download
    finishes the job rather than printing empty cells."""

    async def seed() -> None:
        await store.create(batch_id="unbuilt", country="GB", kind=store.MANUAL, snapshot={})
        await store.add_items(
            "unbuilt",
            [{"service_id": "GB1", "channel_name": "Worst", "country": "GB", "average_pct": 0.9}],
        )
        await store.update_item("unbuilt", "GB1", status="ANALYSED")
        await store.update("unbuilt", status=store.COMPLETED)

    api.portal.call(seed)  # type: ignore[attr-defined]
    before = api.get("/api/batch/batches/unbuilt").json()
    assert before["items"][0]["summary"] is None

    response = api.get("/api/batch/batches/unbuilt/report.csv")
    assert response.status_code == 200, response.text
    assert "The analysis completed and recorded no finding." in response.text

    after = api.get("/api/batch/batches/unbuilt").json()
    assert after["items"][0]["summary"]
    log = api.get("/api/batch/batches/unbuilt/log").json()["lines"]
    assert any("never built" in line["message"] for line in log)
