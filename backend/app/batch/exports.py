"""The automated batch report, as XLSX and CSV.

The column order is fixed by the specification and nothing may reorder it:

    Channel Name | Service ID | Country | Avg Rebuffering Ratio (<window>) |
    Week [start date - end date] | Analysis Summary

The window in the fourth heading is not a constant. It is written from the window the batch
actually measured, read back from the batch row, so a report can never claim seven days while
its averages cover a different span. The fifth column repeats that window as explicit dates,
because a heading is easy to miss and a spreadsheet outlives the screen it came from.

Rows are sorted by average rebuffering ratio, worst first. Channels that failed are listed in
their own section with the reason — never dropped, because a report that quietly omits six
channels is worse than one that names them.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
from typing import Any

REPORT_SHEET = "channels"
FAILED_SHEET = "failed"
CORRELATION_SHEET = "correlation"
ABOUT_SHEET = "about"

FAILED_COLUMNS = ("channel_name", "service_id", "country", "status", "reason")

CORRELATION_COLUMNS = (
    "channel_name",
    "service_id",
    "spike_start_utc",
    "spike_end_utc",
    "peak_rebuffering_pct",
    "spike_minutes",
    "aging_events_in_window",
    "observation",
    "aging_job_id",
)

# Statuses that mean the channel was analysed, so it belongs in the main table.
ANALYSED = ("ANALYSED", "COMPLETED")


def window_label(days: float) -> str:
    """How the fourth column names its window. Seven days reads as "7 days"."""
    whole = round(days)
    if abs(days - whole) < 0.05:
        return f"{whole} day{'s' if whole != 1 else ''}"
    return f"{days:.1f} days"


def headings(window_days: float) -> tuple[str, ...]:
    """The column order, with the window written into the fourth heading."""
    return (
        "Channel Name",
        "Service ID",
        "Country",
        f"Avg Rebuffering Ratio ({window_label(window_days)})",
        "Week [start date - end date]",
        "Analysis Summary",
    )


def _dates(window_from: str | None, window_to: str | None) -> str:
    """The window as explicit dates, with the zone stated."""
    if not window_from or not window_to:
        return ""
    start = dt.datetime.fromisoformat(window_from)
    end = dt.datetime.fromisoformat(window_to)
    return f"{start.strftime('%Y-%m-%d')} - {end.strftime('%Y-%m-%d')} UTC"


def _pct(value: float | None) -> str:
    return "" if value is None else f"{value:.4f}"


def _window_days(batch: dict[str, Any]) -> float:
    if not batch.get("window_from") or not batch.get("window_to"):
        return float((batch.get("settings") or {}).get("window_days", 7))
    start = dt.datetime.fromisoformat(str(batch["window_from"]))
    end = dt.datetime.fromisoformat(str(batch["window_to"]))
    return (end - start).total_seconds() / 86400


def _rows(batch: dict[str, Any]) -> list[list[Any]]:
    """One row per analysed channel, worst first."""
    dates = _dates(batch.get("window_from"), batch.get("window_to"))
    analysed = [item for item in batch.get("items", []) if item["status"] in ANALYSED]
    analysed.sort(key=lambda item: item.get("average_pct") or 0.0, reverse=True)
    return [
        [
            item["channel_name"],
            item["service_id"],
            item["country"],
            _pct(item.get("average_pct")),
            dates,
            item.get("summary") or "",
        ]
        for item in analysed
    ]


def _failed_rows(batch: dict[str, Any]) -> list[list[Any]]:
    """Every channel that was selected but did not produce an analysis."""
    return [
        [
            item["channel_name"],
            item["service_id"],
            item["country"],
            item["status"],
            item.get("error") or item.get("summary") or "",
        ]
        for item in batch.get("items", [])
        if item["status"] not in ANALYSED
    ]


def _correlation_rows(batch: dict[str, Any]) -> list[list[Any]]:
    """Every spike window and what the aging run captured in it."""
    rows: list[list[Any]] = []
    for item in batch.get("items", []):
        for window in (item.get("correlation") or {}).get("windows", []):
            rows.append(
                [
                    item["channel_name"],
                    item["service_id"],
                    window.get("start", ""),
                    window.get("end", ""),
                    _pct(window.get("peak_pct")),
                    window.get("minutes", 0),
                    window.get("event_count", 0),
                    window.get("sentence", ""),
                    item.get("aging_job_id") or "",
                ]
            )
    return rows


def _about(batch: dict[str, Any]) -> list[list[str]]:
    """What produced this report, so the numbers can be read a month later."""
    settings = batch.get("settings") or {}
    return [
        ["field", "value"],
        ["report", "Samsung TV Plus - automated batch"],
        ["country", batch.get("country", "")],
        ["batch_id", batch.get("id", "")],
        ["batch_type", batch.get("kind", "")],
        ["status", batch.get("status", "")],
        ["run_started_utc", batch.get("started_at") or ""],
        ["run_finished_utc", batch.get("finished_at") or ""],
        ["window_start_utc", batch.get("window_from") or ""],
        ["window_end_utc", batch.get("window_to") or ""],
        ["window_days", window_label(_window_days(batch))],
        ["threshold_pct", str(settings.get("threshold_pct", ""))],
        ["analysis_duration_minutes", str(settings.get("analysis_duration_minutes", ""))],
        ["analysis_concurrency", str(settings.get("analysis_concurrency", ""))],
        ["channels_listed", str(batch.get("channels_listed", 0))],
        ["channels_scanned", str(batch.get("channels_scanned", 0))],
        ["channels_above_threshold", str(batch.get("channels_above", 0))],
        ["channels_analysed", str(batch.get("channels_analysed", 0))],
        ["channels_failed", str(batch.get("channels_failed", 0))],
        [
            "selection",
            "Channels whose average rebuffering ratio over the window is above the threshold",
        ],
    ]


def csv_bytes(batch: dict[str, Any]) -> bytes:
    """The report as one CSV: the table first, then the failures and the context."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    writer.writerow(list(headings(_window_days(batch))))
    writer.writerows(_rows(batch))

    failed = _failed_rows(batch)
    if failed:
        writer.writerow([])
        writer.writerow(["Channels selected but not analysed"])
        writer.writerow(list(FAILED_COLUMNS))
        writer.writerows(failed)

    correlated = _correlation_rows(batch)
    if correlated:
        writer.writerow([])
        writer.writerow(["Rebuffering spike windows and what aging captured in them"])
        writer.writerow(list(CORRELATION_COLUMNS))
        writer.writerows(correlated)

    writer.writerow([])
    writer.writerows(_about(batch))
    return buffer.getvalue().encode("utf-8")


def xlsx_bytes(batch: dict[str, Any]) -> bytes:
    """The report as a workbook: a sheet per section, each a table with its header on row 1."""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = REPORT_SHEET
    sheet.append(list(headings(_window_days(batch))))
    for row in _rows(batch):
        sheet.append(row)
    sheet.freeze_panes = "A2"

    failed = _failed_rows(batch)
    if failed:
        page = workbook.create_sheet(FAILED_SHEET)
        page.append(list(FAILED_COLUMNS))
        for row in failed:
            page.append(row)
        page.freeze_panes = "A2"

    correlated = _correlation_rows(batch)
    if correlated:
        page = workbook.create_sheet(CORRELATION_SHEET)
        page.append(list(CORRELATION_COLUMNS))
        for row in correlated:
            page.append(row)
        page.freeze_panes = "A2"

    about = workbook.create_sheet(ABOUT_SHEET)
    for line in _about(batch):
        about.append(line)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def filename(country: str, fmt: str, now: dt.datetime | None = None) -> str:
    stamp = (now or dt.datetime.now(dt.UTC)).strftime("%Y%m%d")
    return f"automated_batch_{country.upper()}_{stamp}.{fmt}"
